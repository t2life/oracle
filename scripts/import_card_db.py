from __future__ import annotations

"""カードDB取込パイプライン。

単一真実源: C:\\Users\\81907\\Desktop\\オラクルカード\\DB\\ のxlsx群
  - 日本神話オラクルカード 44柱.xlsx（製品マスタ: No.1〜44・テーマ別解釈3系統）
  - 日本神話オラクルカード DB.xlsx（AI解釈用キーワード辞書・組み合わせ解釈ルール・
    文脈パターンDB・感情トーン設定・質問タイプ分類）

出力: src/oracle_app/data/card_content.json（手動編集禁止）
xlsx更新時は本スクリプトを再実行し、export_master_assets.py→モバイル再ビルドまで行うこと。

注記（2026-09-06実測）: カードデータベース.xlsxの「アプリ実装マスタ」は29件かつ
神名表記が別体系（例: 素戔嗚尊 vs 須佐之男命）で44柱と結合不能のため取込対象外。
英語名はv1では未収録（UIは日本語神名へフォールバック）。
"""

import json
import sys
from pathlib import Path

import openpyxl

ROOT_DIR = Path(__file__).resolve().parents[1]
DB_DIR = Path(r"C:\Users\81907\Desktop\オラクルカード\DB")
OUTPUT_PATH = ROOT_DIR / "src" / "oracle_app" / "data" / "card_content.json"

CARD_COUNT = 44


def _split_words(value: object) -> list[str]:
    if value is None:
        return []
    text = str(value).replace("，", "、").replace(",", "、")
    return [word.strip() for word in text.split("、") if word.strip()]


def _rows(ws, min_row: int = 2) -> list[tuple]:
    return [
        row
        for row in ws.iter_rows(min_row=min_row, values_only=True)
        if any(cell not in (None, "") for cell in row)
    ]


def main() -> None:
    wb44 = openpyxl.load_workbook(
        DB_DIR / "日本神話オラクルカード 44柱.xlsx", data_only=True
    )
    wbdb = openpyxl.load_workbook(
        DB_DIR / "日本神話オラクルカード DB.xlsx", data_only=True
    )

    base_rows = _rows(wb44["44柱一覧"])
    if len(base_rows) != CARD_COUNT:
        raise SystemExit(f"44柱一覧の件数が{len(base_rows)}件です（44件を期待）。")

    theme_tabs = {
        "love": _rows(wb44["恋愛リーディング解釈"]),
        "work_money": _rows(wb44["仕事・金運リーディング解釈"]),
        "health_mind": _rows(wb44["健康・精神リーディング解釈"]),
    }
    for label, rows in theme_tabs.items():
        if len(rows) != CARD_COUNT:
            raise SystemExit(f"{label}タブの件数が{len(rows)}件です（44件を期待）。")

    keyword_dict: dict[str, dict[str, list[str]]] = {}
    for row in _rows(wbdb["AI解釈用キーワード辞書"]):
        name = str(row[1]).strip()
        keyword_dict[name] = {
            "positive": _split_words(row[2]),
            "negative": _split_words(row[3]),
            "neutral": _split_words(row[4]),
            "action": _split_words(row[5]),
        }

    cards = []
    for index, row in enumerate(base_rows):
        no = int(row[0])
        name = str(row[1]).strip()
        love_row = theme_tabs["love"][index]
        work_row = theme_tabs["work_money"][index]
        health_row = theme_tabs["health_mind"][index]
        if str(love_row[1]).strip() != name:
            raise SystemExit(f"No.{no} 神名不一致: 一覧={name} 恋愛={love_row[1]}")

        love_text = str(love_row[2]).strip()
        work_text = str(work_row[2]).strip()
        health_text = str(health_row[2]).strip()
        dictionary = keyword_dict.get(name)
        if dictionary is None:
            raise SystemExit(f"No.{no} {name} がAI解釈用キーワード辞書にありません。")

        cards.append(
            {
                "no": no,
                "card_id": f"japanese_mythology_card_{no:03d}",
                "name_ja": name,
                "reading": str(row[2] or "").strip(),
                "attribute": str(row[3] or "").strip(),
                "element": str(row[4] or "").strip(),
                "basic_meaning": str(row[5] or "").strip(),
                "keywords": _split_words(row[6]),
                # テーマ別解釈: 仕事・金運タブ→work/money、健康・精神タブ→health/潜在意識/ハイヤーセルフ
                "theme_meanings": {
                    "love": love_text,
                    "work": work_text,
                    "money": work_text,
                    "health": health_text,
                    "subconscious": health_text,
                    "higher_self": health_text,
                },
                "keyword_dict": dictionary,
            }
        )

    context_patterns = [
        {
            "genre": str(row[0]).strip(),
            "situation": str(row[1]).strip(),
            "templates": [
                str(cell).strip()
                for cell in (row[2], row[3], row[4])
                if cell not in (None, "")
            ],
        }
        for row in _rows(wbdb["文脈パターンDB"])
    ]

    tones = [
        {
            "name": str(row[0]).strip(),
            "scene": str(row[1] or "").strip(),
            "style": str(row[2] or "").strip(),
            "endings": _split_words(row[3]),
            "avoid": _split_words(row[4]),
        }
        for row in _rows(wbdb["感情トーン設定"])
    ]

    # 質問タイプ分類は 2026-09-10 に判定用の列を追加した（状況判定1/2・優先度）。
    # 旧レイアウトのxlsxでも動くよう、列が無い場合は既存列から埋める（fail-soft）。
    def _cell(row: tuple, index: int) -> str:
        return str(row[index] or "").strip() if len(row) > index else ""

    question_types = [
        {
            "pattern": str(row[0]).strip(),
            "genre": _cell(row, 1),
            "situation": _cell(row, 2),
            "tone": _cell(row, 3),
            # 相談内容の判定に使う語（読点区切り）と、複数候補に分解した状況
            "keywords": _split_words(row[4] if len(row) > 4 else None),
            "situations": [
                value
                for value in (_cell(row, 5) or _cell(row, 2), _cell(row, 6))
                if value
            ],
            "priority": row[7] if len(row) > 7 and row[7] is not None else 999,
        }
        for row in _rows(wbdb["質問タイプ分類"])
    ]

    spreads = [
        {
            "spread_id": str(row[0]).strip(),
            "name_ja": _cell(row, 1),
            "kind": _cell(row, 2),
            "card_count": int(row[3] or 0),
            "min_cards": int(row[4] or 1),
            "max_cards": int(row[5] or 1),
            "purpose": _cell(row, 6),
            "required_tickets": int(row[7] or 0),
            "allowed_plans": _split_words(row[8]),
            "sort_order": int(row[9] or 0),
        }
        for row in _rows(wbdb["スプレッド定義"])
    ]

    positions: dict[str, list[dict[str, object]]] = {}
    for row in _rows(wbdb["ポジション定義"]):
        positions.setdefault(str(row[0]).strip(), []).append(
            {
                "index": int(row[1] or 0),
                "name": _cell(row, 2),
                "meaning": _cell(row, 3),
                "role": _cell(row, 4),
            }
        )
    for items in positions.values():
        items.sort(key=lambda item: item["index"])

    connectors = [
        {
            "tone": str(row[0]).strip(),
            "opening": _cell(row, 1),
            "linking": _cell(row, 2),
            "closing": _cell(row, 3),
        }
        for row in _rows(wbdb["接続表現マスタ"])
    ]

    theme_genres = {
        str(row[0]).strip(): _cell(row, 2)
        for row in _rows(wbdb["テーマジャンル対応"])
    }

    combination_rules = [
        {
            "attr1": str(row[0]).strip(),
            "attr2": str(row[1]).strip(),
            "interaction": str(row[2] or "").strip(),
            "direction": str(row[3] or "").strip(),
            "keywords": _split_words(row[4]),
        }
        for row in _rows(wbdb["組み合わせ解釈ルール"])
    ]

    payload = {
        "_generated_by": "scripts/import_card_db.py",
        "_note": "手動編集禁止。単一真実源はDBフォルダのxlsx。更新時は本スクリプト→export_master_assets.py→モバイル再ビルド。",
        "cards": cards,
        "context_patterns": context_patterns,
        "tones": tones,
        "question_types": question_types,
        "combination_rules": combination_rules,
        "spreads": spreads,
        "positions": positions,
        "connectors": connectors,
        "theme_genres": theme_genres,
    }

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(
        json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    size_kb = OUTPUT_PATH.stat().st_size / 1024
    print(f"カードコンテンツを書き出しました: {OUTPUT_PATH} ({size_kb:.0f} KB)")
    print(
        f"カード{len(cards)}件 / 文脈{len(context_patterns)}件 / トーン{len(tones)}件 / "
        f"質問タイプ{len(question_types)}件 / 組み合わせ{len(combination_rules)}件 / "
        f"スプレッド{len(spreads)}件 / ポジション{sum(len(v) for v in positions.values())}件 / "
        f"接続表現{len(connectors)}件 / テーマ対応{len(theme_genres)}件"
    )


if __name__ == "__main__":
    sys.exit(main())
